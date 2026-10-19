"""Tabular uploads: converted once to Parquet, read by DuckDB where they lie.

CHUNKING A CSV IS A CATEGORY ERROR, and it is the default in most RAG
pipelines. Two hundred rows of a price table become two hundred near-identical
embeddings that rank against each other and answer nothing -- ask "what was
August revenue" and you retrieve five rows that all look equally like the
question. A table's query language already exists.

So a tabular upload takes the other path: convert to Parquet once (columnar,
typed, compressed, and the format every engine reads), leave it in object
storage next to the original bytes, and register it in `content.datasets` so
DuckDB can read it straight from S3.

HOW ICEBERG-LIKE THIS IS, precisely, because the phrase invites more than is
here. Files for one dataset land under one prefix, and `mode: append` means the
reader globs the prefix -- so this month's export joins last month's and
`select * from read_parquet('s3://…/sales/*.parquet')` sees both. That is
enough to be useful and it is all of it: no manifest, no snapshot isolation, no
schema evolution, no time travel, no compaction, no deletes. What breaks first
is a schema change between two files under one prefix, which DuckDB will union
by name or refuse depending on how far apart they are. Iceberg exists because
those are hard; when a deployment needs them it needs Iceberg, and the
migration is `read_parquet` to a table, not a rewrite of this module.
"""

from __future__ import annotations

import io
import posixpath
import re

from .policy import IngestPolicy

__all__ = ["dataset_name_for", "object_key_for", "to_parquet"]


class MissingParser(RuntimeError):
    """pyarrow is not installed."""


def _arrow():
    try:
        import pyarrow  # noqa: F401
    except ModuleNotFoundError as e:
        raise MissingParser(
            "tabular ingestion needs pyarrow: pip install 'percolate-core[tabular]' "
            f"(missing {e.name})") from None
    import pyarrow as pa
    return pa


def _read_csv(data: bytes, policy: IngestPolicy, delimiter: str | None):
    _arrow()   # the import check, with its message; pyarrow.csv needs it loaded
    from pyarrow import csv as pacsv

    parse = pacsv.ParseOptions(delimiter=delimiter) if delimiter else pacsv.ParseOptions()
    read = pacsv.ReadOptions(
        # A file with no header gets f0, f1, … rather than losing its first row
        # of data to a header it never had.
        autogenerate_column_names=not policy.table.header)
    return pacsv.read_csv(io.BytesIO(data), read_options=read, parse_options=parse)


def _read_xlsx(data: bytes, policy: IngestPolicy):
    """openpyxl if it is there, and a clear refusal if it is not.

    Not in the extra: openpyxl is small, but a spreadsheet is a document with
    several sheets, merged cells and formulas, and treating sheet one as "the
    table" is a guess. Deployments that upload spreadsheets should say so.
    """
    pa = _arrow()
    try:
        import openpyxl
    except ModuleNotFoundError:
        raise MissingParser(
            "reading .xlsx needs openpyxl (pip install openpyxl). CSV and Parquet "
            "need nothing beyond the tabular extra.") from None

    book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = book[book.sheetnames[0]]
    rows = [[c for c in row] for row in sheet.iter_rows(values_only=True)]
    if not rows:
        raise ValueError("the first sheet is empty")
    header = [str(h) if h is not None else "f" + str(i)
              for i, h in enumerate(rows[0])] if policy.table.header else \
             ["f" + str(i) for i in range(len(rows[0]))]
    body = rows[1:] if policy.table.header else rows
    cols = {name: [r[i] if i < len(r) else None for r in body]
            for i, name in enumerate(header)}
    return pa.table(cols)


def to_parquet(data: bytes, policy: IngestPolicy,
               name: str | None = None) -> tuple[bytes, list[dict[str, str]], int]:
    """(parquet bytes, columns, row count).

    Parquet in, Parquet out is still a re-encode rather than a passthrough, and
    that is on purpose: it is how the registry's column list comes from a
    schema this process actually read rather than from a claim about the file.
    """
    _arrow()
    import pyarrow.parquet as pq

    ext = posixpath.splitext((name or "").lower())[1]
    if ext == ".parquet":
        table = pq.read_table(io.BytesIO(data))
    elif ext in (".xlsx", ".xls"):
        table = _read_xlsx(data, policy)
    else:
        delimiter = policy.table.delimiter or ("\t" if ext == ".tsv" else None)
        table = _read_csv(data, policy, delimiter)

    buf = io.BytesIO()
    pq.write_table(table, buf, compression="zstd")
    columns = [{"name": f.name, "type": str(f.type)} for f in table.schema]
    return buf.getvalue(), columns, table.num_rows


def dataset_name_for(policy: IngestPolicy, name: str | None, title: str | None) -> str:
    """The logical table's name: the policy's, else the file's stem.

    Slugged, because it becomes a path segment and a DuckDB identifier. Never
    guessed across files -- `sales-2026-08.csv` is a dataset called
    `sales-2026-08` unless the uploader says `table: {name: sales}`. Stripping
    a date suffix to merge two uploads is the kind of helpfulness that silently
    unions two unrelated exports.
    """
    if policy.table.name:
        raw = policy.table.name
    else:
        raw = posixpath.splitext(posixpath.basename(name or title or "dataset"))[0]
    slug = re.sub(r"[^a-zA-Z0-9_-]+", "-", raw).strip("-").lower()
    return slug or "dataset"


def object_key_for(dataset: str, checksum: str) -> str:
    """One prefix per dataset, one object per upload, named by content hash.

    The hash rather than a sequence number, for the same reason `objects/` is
    content-addressed: re-ingesting the same file overwrites itself instead of
    appending a duplicate that the glob would then count twice.
    """
    return "datasets/" + dataset + "/" + checksum + ".parquet"


# `read_dataset` WAS HERE, and it is deleted rather than fixed (REM-69).
#
# It opened its own DuckDB connection, ran `install httpfs` at query time, and
# carried the third of this package's three different readings of the S3
# environment. It also had no callers -- not in this package, not in the tests,
# not in the UI. A second, unsealed way to reach object storage, maintained for
# nobody, is worse than none: `challenge-the-process` §3, deleting a step beats
# automating one.
#
# The capability is not lost. `lake/engine.py` is the module that reaches object
# storage from DuckDB and is the one under test; a caller that needs to query a
# dataset where it lies should go through it, and if it turns out to want a
# convenience wrapper, that wrapper belongs there beside the connection it uses.

